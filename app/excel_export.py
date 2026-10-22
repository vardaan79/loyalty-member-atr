"""V2-house-format Excel export for the Loyalty ATR app.

Uses the locked house module E:\\CRM_Reports\\scripts\\report_format_v2.py:
title (merged) -> header band R4 -> GRD-TOT R6 / SUB-TOT R7 (SUM / SUBTOTAL,
sales values total in LAKHS per house convention) -> spacer R8 -> data R9+,
freeze C9, zoom 85, gridlines on, yellow spacer column between label and measures.
"""
import os, sys, datetime
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, r'E:\CRM_Reports\scripts')
import report_format_v2 as V2

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
os.makedirs(OUT, exist_ok=True)

MAX_LIST_ROWS = 5000          # campaign lists capped in Excel (CSVs carry the full ranked list)

# column kinds
K = "key"; T = "text"; I = "int"; M = "money"; P = "pct"; D = "dec"; DT = "date"; AM = "money-nototal"
# kind -> data-zone number format (mobiles keep leading zeros as text)
ZONE_FMT = {K: "@", T: "General", I: V2.NF_INT, M: V2.NF_ACCT, P: V2.NF_PCT1,
            D: V2.NF_DEC, DT: V2.NF_DATE, AM: V2.NF_ACCT}
TIER_PRICE = {1: 200, 2: 500, 3: 1000}
TIER_RATE = {1: 0.02, 2: 0.05, 3: 0.11}


def _pct(v):
    """Store fractions so NF_PCT renders them; None-safe."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    return float(v) / 100.0


def _date(v):
    if v is None or pd.isna(v):
        return None
    return pd.to_datetime(v).date()


def _sheet(wb, name, title, cols, rows, plain_sum=()):
    """One canonical V2 sheet. cols = [(header, kind, width)]; rows = list of tuples.
    plain_sum: 1-based col indices (of `cols`) totalled as plain SUM (counts)."""
    ws = wb.create_sheet(name[:31])
    last_col = len(cols) + 2                      # A label + B spacer + data cols
    V2.title(ws, f"C1:{get_column_letter(last_col)}1", title)
    ws.cell(row=2, column=3,
            value=f"Generated {datetime.datetime.now().strftime('%d-%b-%y %H:%M')} · data window per cover").font = V2.F_NOTE
    for i, (h, kind, w) in enumerate(cols):
        V2.header(ws, 4, i + 3, h, kind=("key" if kind == K else "default"))
        ws.column_dimensions[get_column_letter(i + 3)].width = w
    ws.cell(row=6, column=1, value="GRD-TOT").font = V2.F_TOTAL
    ws.cell(row=7, column=1, value="SUB-TOT").font = V2.F_TOTAL
    # totals: money sums in LAKHS (house rule); counts via plain_sum override; rest skipped
    kinds = {}
    for i, (_, kind, _) in enumerate(cols):
        col = i + 3
        if col - 2 in plain_sum:                  # plain_sum given in `cols` positions
            kinds[col] = "__plain"
        else:
            kinds[col] = "money" if kind == M else "text"
    V2.total_rows(ws, 1, last_col, kinds={c: ("text" if v != "money" else v) for c, v in kinds.items()})
    for col, v in kinds.items():                  # rewrite plain-SUM totals after house rows
        if v == "__plain":
            L = get_column_letter(col)
            for r in (6, 7):
                c = ws.cell(row=r, column=col, value=f"=SUM({L}9:{L}1048576)")
                c.font, c.border, c.number_format = V2.F_TOTAL, V2.B_ALL, V2.NF_INT
    V2.style_data_zone(ws, 1, last_col,
                       numfmts={i + 3: ZONE_FMT[k] for i, (_, k, _) in enumerate(cols)},
                       n_rows=max(len(rows), 1))
    V2.spacer_column(ws, 2, first_row=4, last_row=max(9, 8 + len(rows)))
    for r, row in enumerate(rows, start=9):
        for i, v in enumerate(row):
            ws.cell(row=r, column=i + 3, value=v)
    ws.column_dimensions["A"].width = 16
    V2.finalize(ws, freeze="C9")


def _campaign_frames(res):
    """Ranked campaign lists from the members frame (same rules as engine.build)."""
    JUNK = {'6666666666', '7777777777', '8888888888', '9999999999', '9876543210'}
    m = res["members"].copy()
    m = m[~m.mobile.isin(JUNK) & m.mobile.str.match(r'^[6-9][0-9]{9}$', na=False)]
    asof = pd.to_datetime(res["to"])
    winback = m[(m.mem_bills >= 2) & (m.last_dt <= asof - datetime.timedelta(days=30))].sort_values("mem_sales", ascending=False)
    onboard = m[m.mem_bills <= 1].sort_values("card_paid", ascending=False)
    upgrade = m[((m.tier == "Silver") & (m.mem_sales >= 25000)) |
                ((m.tier == "Gold") & (m.mem_sales >= 50000))].sort_values("mem_sales", ascending=False)
    med = m.loc[m.mem_sales > 0, "mem_sales"].median()
    trade = m[(m.mem_sales >= max(10000, 2 * (med or 0))) &
              (m.mem_ev / m.mem_sales.replace(0, pd.NA) >= 0.8)].sort_values("mem_sales", ascending=False)
    return dict(winback=winback, onboard=onboard, upgrade=upgrade, tradeup=trade)


def export(res, path=None):
    """Build the full workbook from an engine results dict. Returns file path."""
    k = res["kpis"]
    frm, to = res["frm"] or "All-time", res["to"]
    stamp = pd.to_datetime(to).strftime("%d-%b-%y")
    sub = f"{frm} to {to}"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    money = lambda v: float(v) if pd.notna(v) else None
    pctf = _pct
    delta = lambda a, b: (a / b - 1) if b else None

    # ---- 01 ATR summary -----------------------------------------------------
    _sheet(wb, "01-ATR-SUMMARY", f"Loyalty Member ATR · {sub}", [
        ("AREA", T, 18), ("OBSERVATION", T, 80), ("RECOMMENDED ACTION", T, 60),
        ("OWNER", T, 14), ("STATUS", T, 14)],
        [(a, o, x, "", "") for a, o, x in res["insights"]])

    # ---- 02 KPI -------------------------------------------------------------
    kp = k
    _sheet(wb, "02-KPI", f"Loyalty Member KPI · {sub}", [
        ("METRIC", T, 22), ("MEMBER", D, 16), ("NON-MEMBER", D, 16), ("DELTA %", P, 12)],
        [("Customers", kp["members"], kp["non_customers"], pctf(delta(kp["members"], kp["non_customers"]))),
         ("Sales value", money(kp["mem_sales"]), money(kp["non_sales"]), pctf(delta(kp["mem_sales"], kp["non_sales"]))),
         ("Bills", kp["mem_bills"], kp["non_bills"], pctf(delta(kp["mem_bills"], kp["non_bills"]))),
         ("Avg bill value", money(kp["abv_m"]), money(kp["abv_n"]), pctf(kp["abv_m"] / kp["abv_n"] - 1)),
         ("Items per bill", round(kp["ipb_m"], 2), round(kp["ipb_n"], 2), pctf(kp["ipb_m"] / kp["ipb_n"] - 1)),
         ("Avg selling price", money(kp["asp_m"]), money(kp["asp_n"]), pctf(kp["asp_m"] / kp["asp_n"] - 1)),
         ("Bills per customer", round(kp["freq_m"], 2), round(kp["freq_n"], 2),
          pctf(kp["freq_m"] / max(kp["freq_n"], 1e-9) - 1)),
         ("GM% (costed lines)", pctf(kp["mem_gmp"]), pctf(kp["non_gmp"]),
          pctf((kp["mem_gmp"] or 0) - (kp["non_gmp"] or 0))),
         ("Member share of sales", pctf(kp["mem_share_pct"]), None, None)])

    # ---- 03 weekly ----------------------------------------------------------
    et = res["enrol_tier"].pivot_table(index="wk", columns="rnk", values="n", aggfunc="sum") if len(res["enrol_tier"]) else res["enrol_tier"]
    enrol_wk = res["enrol"].set_index("wk")["new_members"].to_dict() if len(res["enrol"]) else {}
    wrows = []
    for _, r in res["weekly"].iterrows():
        wk = pd.to_datetime(r.wk)
        g = lambda rk: int(et.loc[wk, rk]) if len(et) and wk in et.index and rk in et.columns and pd.notna(et.loc[wk, rk]) else 0
        wrows.append((_date(wk), money(r.mem_v), money(r.tot_v), pctf(r.mem_v / r.tot_v if r.tot_v else None),
                      int(enrol_wk.get(wk, 0)), g(1), g(2), g(3)))
    _sheet(wb, "03-WEEKLY", f"Loyalty Member Weekly Trend · {sub}", [
        ("WEEK\nSTARTING", DT, 12), ("MEMBER\nSALES", M, 14), ("TOTAL\nSALES", M, 14),
        ("MEMBER\nSHARE", P, 10), ("NEW\nMEMBERS", I, 10),
        ("NEW\nSILVER", I, 9), ("NEW\nGOLD", I, 9), ("NEW\nPLATINUM", I, 10)], wrows,
        plain_sum=(5, 6, 7, 8))

    # ---- mix sheets (04-09) -------------------------------------------------
    def mix_rows(df, label_cols, cap=None):
        rows = []
        for _, r in (df.head(cap) if cap else df).iterrows():
            rows.append(tuple([r[c] for c in label_cols] + [
                money(r.mem_v), int(r.mem_q), int(r.mem_bills),
                pctf(r.mem_share), pctf(r.non_share),
                int(round(r.ix)) if pd.notna(r.ix) else None,
                pctf(r.e_v / r.mem_v) if r.mem_v else None,
                pctf(r.v_v / r.mem_v) if r.mem_v else None,
                pctf(r.p_v / r.mem_v) if r.mem_v else None,
                pctf(r.sp_v / r.mem_v) if r.mem_v else None]))
        return rows

    MIX_COLS = [("MEMBER\nSALES", M, 13), ("QTY", I, 9), ("MEM\nBILLS", I, 9),
                ("MEM\nSHARE", P, 9), ("NON\nSHARE", P, 9), ("INDEX", I, 8),
                ("E %", P, 8), ("V %", P, 8), ("P %", P, 8), ("SP %", P, 8)]

    _sheet(wb, "04-DIVISION", f"Mix · Division · {sub}",
           [("DIVISION", K, 12)] + MIX_COLS, mix_rows(res["div"], ["division"]))
    _sheet(wb, "05-SUB-DIVISION", f"Mix · Sub Division (Top 60) · {sub}",
           [("DIVISION", T, 10), ("SUB\nDIVISION", K, 18)] + MIX_COLS,
           mix_rows(res["subdiv"], ["division", "sub_division"], 60))
    _sheet(wb, "06-MAJ-CATEGORY", f"Mix · Major Category (Top 60) · {sub}",
           [("DIVISION", T, 10), ("MAJOR\nCATEGORY", K, 18)] + MIX_COLS,
           mix_rows(res["majcat"], ["division", "maj_cat"], 60))
    _sheet(wb, "07-PRICE-GRID", f"Mix · Price Grid E/V/P/SP · {sub}",
           [("GRID", K, 8)] + MIX_COLS,
           mix_rows(res["grid"][res["grid"].grid.isin(["E", "V", "P", "SP"])], ["grid"]))
    _sheet(wb, "08-SIZE", f"Mix · Size (Top 40) · {sub}",
           [("DIVISION", T, 10), ("SIZE", K, 10)] + MIX_COLS,
           mix_rows(res["size"][res["size"]["size"] != "NA"], ["division", "size"], 40))
    _sheet(wb, "09-COLOUR", f"Mix · Colour (Top 30) · {sub}",
           [("COLOUR", K, 14)] + MIX_COLS,
           mix_rows(res["colour"][res["colour"].colour != "NA"], ["colour"], 30))

    # ---- 10 tiers -----------------------------------------------------------
    RANK = {"Silver": 1, "Gold": 2, "Platinum": 3}
    trows = []
    for t, r in res["tier_tbl"].iterrows():
        trows.append((t, f"₹{TIER_PRICE[RANK[t]]:,} · {TIER_RATE[RANK[t]]*100:.0f}% slab",
                      int(r.members), pctf(r.members / max(res['tier_tbl'].members.sum(), 1)),
                      money(r.mem_sales), pctf(r.mem_sales / max(res['tier_tbl'].mem_sales.sum(), 1)),
                      money(r.card_rev), int(r.buyers), round(r.freq, 2),
                      money(r.mem_abv) if pd.notna(r.mem_abv) else None,
                      pctf(r.gmp), money(r.slab_disc),
                      pctf(r.e_pct), pctf(r.v_pct), pctf(r.p_pct), pctf(r.sp_pct)))
    _sheet(wb, "10-TIERS", f"Tier Segmentation (Silver / Gold / Platinum) · {sub}", [
        ("TIER", K, 10), ("CARD", T, 16), ("MEMBERS", I, 10), ("MEMBER\nSHARE", P, 10),
        ("MEMBER\nSALES", M, 14), ("SALES\nSHARE", P, 10), ("CARD\nREVENUE", M, 12),
        ("CARD\nBUYERS", I, 9), ("BILLS /\nMEMBER", D, 9), ("ABV", M, 12), ("GM %", P, 9),
        ("EST. SLAB\nDISC", M, 12), ("E %", P, 8), ("V %", P, 8), ("P %", P, 8), ("SP %", P, 8)],
        trows, plain_sum=(3, 4, 8))

    # ---- 11 employees -------------------------------------------------------
    dd_bills = res["emp_dd"].groupby("mobile").transactionid.nunique().to_dict()
    dd_amt = res["emp_dd"].groupby("mobile").empdisc.sum().to_dict()
    emp = res["emp_mem"].sort_values("mem_sales", ascending=False).head(MAX_LIST_ROWS)
    erows = [(str(r.ecode), str(r.name), str(r.desig), str(r.loc), r.mobile, r.tier, r.status,
              int(r.card_bills), money(r.card_paid), int(r.mem_bills), money(r.mem_sales),
              int(dd_bills.get(r.mobile, 0)), money(dd_amt.get(r.mobile, 0)))
             for r in emp.itertuples()]
    _sheet(wb, "11-EMPLOYEES", f"Employee Members (Top {len(erows):,} by Spend) · {sub}", [
        ("E-CODE", K, 10), ("NAME", T, 24), ("DESIGNATION", T, 18), ("LOCATION", T, 16),
        ("MOBILE", K, 12), ("TIER", T, 9), ("STATUS", T, 10), ("CARD\nBILLS", I, 8), ("CARD\nPAID", M, 11),
        ("MEMBER\nBILLS", I, 9), ("MEMBER\nSALES", M, 13), ("DD\nBILLS", I, 8), ("DD\nAMOUNT", M, 11)],
        erows, plain_sum=(8, 10, 12))

    # ---- 12 GM by division (+ tier split) ------------------------------------
    gdt = res.get("gm_div_tier")
    gdt = gdt.set_index("division") if gdt is not None and len(gdt) else None
    def _tg(d, k):
        if gdt is None or d not in gdt.index:
            return None
        v = gdt.loc[d, k]
        return pctf(v) if v is not None and pd.notna(v) else None
    grows = []
    for r in res["gm_div"].itertuples():
        grows.append((r.division, money(r.mem_v),
                      pctf(100 * r.mem_gm / r.mem_costed) if r.mem_costed else None,
                      _tg(r.division, "s_gmp"), _tg(r.division, "g_gmp"), _tg(r.division, "p_gmp"),
                      money(r.non_v),
                      pctf(100 * r.non_gm / r.non_costed) if r.non_costed else None))
    _sheet(wb, "12-GM-DIVISION", f"Member GM by Division with Tier Split (Costed Lines) · {sub}", [
        ("DIVISION", K, 12), ("MEMBER\nSALES", M, 14), ("MEMBER\nGM %", P, 10),
        ("SILVER\nGM %", P, 10), ("GOLD\nGM %", P, 10), ("PLATINUM\nGM %", P, 10),
        ("NON-MEMBER\nSALES", M, 14), ("NON-MEMBER\nGM %", P, 10)], grows)

    # ---- 13 campaign plan + 14-18 lists -------------------------------------
    CAM_META = {
        "winback":  ("Dormant winback", "2+ member bills, no purchase in 30+ days", "RCS + call-centre, tier-tagged offer"),
        "onboard":  ("Onboarding nudge", "Bought a card, never returned (1 bill)", "Welcome journey SMS/RC with ₹-off next-visit coupon"),
        "upgrade":  ("Card upgrade", "Silver ≥ ₹25k / Gold ≥ ₹50k member-period spend", "Next-tier slab pitch (5%/11%) vs card fee"),
        "tradeup":  ("Premium trade-up", "≥80% of member spend in E/V grids", "P/SP new-arrival alerts + curated premium edits"),
        "suppress": ("Employee suppression", "Employee mobiles holding cards", "Exclude from all member comms; audit double-dip"),
    }
    crows = [(nm, rule, res["campaigns"][key]["n"], money(res["campaigns"][key]["value"]), chan,
              res["campaigns"][key]["file"])
             for key, (nm, rule, chan) in CAM_META.items()]
    _sheet(wb, "13-CAMPAIGNS", f"CRM Campaign Plan · {sub}", [
        ("CAMPAIGN", K, 18), ("RULE", T, 42), ("TARGETS", I, 10),
        ("MEMBER-PERIOD\nSALES", M, 15), ("CHANNEL / OFFER", T, 40), ("CSV FILE", T, 34)],
        crows, plain_sum=(3,))

    lists = _campaign_frames(res)
    LIST_COLS = [("MOBILE", K, 12), ("TIER", T, 9), ("CARD\nSINCE", DT, 11), ("CARD\nPAID", M, 11),
                 ("MEM\nBILLS", I, 8), ("MEM\nSALES", M, 13), ("LAST\nSHOPPED", DT, 11), ("TOP\nSTORE", T, 8)]
    for key, (sheet_nm, ttl) in {"winback": ("14-WINBACK", "Winback · Dormant Members"),
                                 "onboard": ("15-ONBOARD", "Onboarding · Single-Bill Members"),
                                 "upgrade": ("16-UPGRADE", "Card Upgrade · High Spend"),
                                 "tradeup": ("17-TRADEUP", "Premium Trade-Up · E/V Heavy")}.items():
        rows = [(r.mobile, r.tier, _date(r.card_since), money(r.card_paid), int(r.mem_bills),
                 money(r.mem_sales), _date(r.last_dt), str(r.store or ""))
                for r in lists[key].head(MAX_LIST_ROWS).itertuples()]
        _sheet(wb, sheet_nm, f"{ttl} (Top {len(rows):,}) · {sub}", LIST_COLS, rows, plain_sum=(5,))

    sup = res["emp_mem"].sort_values("mem_sales", ascending=False)
    rows = [(r.mobile, r.tier, _date(r.card_since), money(r.card_paid), int(r.mem_bills),
             money(r.mem_sales), r.status, str(r.ecode), str(r.name), str(r.desig), str(r.loc))
            for r in sup.head(MAX_LIST_ROWS).itertuples()]
    _sheet(wb, "18-SUPPRESS", f"Employee Suppression List · {sub}", [
        ("MOBILE", K, 12), ("TIER", T, 9), ("CARD\nSINCE", DT, 11), ("CARD\nPAID", M, 11),
        ("MEM\nBILLS", I, 8), ("MEM\nSALES", M, 13), ("STATUS", T, 10), ("E-CODE", K, 10),
        ("NAME", T, 24), ("DESIGNATION", T, 18), ("LOCATION", T, 16)], rows, plain_sum=(5,))

    # ---- 19 customers: BEFORE card-day / card-day / AFTER (full member registry) ----
    cu = res.get("customers")
    if cu is not None:
        crows = [(r.mobile, r.tier, str(r.card_store or ""), _date(r.card_since), int(r.days_since),
                  int(r.bef_visits_all), _date(r.bef_last), int(r.bef_visits_12m), int(r.bef_bills_12m),
                  money(r.bef_val_12m), money(r.bef_abv_12m) if pd.notna(r.bef_abv_12m) else None,
                  int(r.cd_bills), money(r.cd_val), money(r.cd_abv) if pd.notna(r.cd_abv) else None,
                  money(r.cd_disc), int(r.aft_visits), int(r.aft_bills), money(r.aft_val),
                  money(r.aft_abv) if pd.notna(r.aft_abv) else None, money(r.aft_disc),
                  money(r.ttl_disc), money(r.card_fee), pctf(r.disc_pct_fee), r.fee_recovered,
                  r.is_emp, r.status, r.high_disc, r.rare_shopper, r.new_cust, r.flag)
                 for r in cu.itertuples()]
        _sheet(wb, "19-CUSTOMERS", f"Customer BEFORE Card-Day / AFTER · {sub}", [
            ("CUST\nMOB", K, 12), ("CARD\nTYPE", T, 9), ("ST-CD", T, 8), ("CARD\nPUR DT", DT, 11),
            ("DAYS\nSINCE", I, 8), ("BEF\nVISITS", I, 8), ("BEF LAST\nVISIT DT", DT, 11),
            ("BEF VIS\n12M", I, 8), ("BEF BILLS\n12M", I, 9), ("BEF BILL VAL\n12M", M, 12),
            ("BEF ABV\n12M", AM, 10), ("CD\nBILLS", I, 8), ("CD BILL\nVAL", M, 12), ("CD ABV", AM, 10),
            ("CD\nDISC", M, 11), ("AFT\nVISITS", I, 8), ("AFT\nBILLS", I, 8), ("AFT BILL\nVAL", M, 12),
            ("AFT ABV", AM, 10), ("AFT\nDISC", M, 11), ("TTL\nDISC", M, 11), ("CARD\nFEE", M, 10),
            ("DISC %\nFEE", P, 9), ("FEE\nRECOVERED", T, 11), ("EMP", T, 6), ("STATUS", T, 10),
            ("HIGH\nDISC", T, 7), ("RARE\nSHOPPER", T, 8), ("NEW\nCUST", T, 7), ("FLAG", T, 30)],
            crows, plain_sum=(6, 8, 9, 12, 16, 17))

    if path is None:
        path = os.path.join(OUT, f"Loyalty_ATR_{str(frm).replace('-', '')}_{to.replace('-', '')}.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    # self-test with a tiny synthetic results dict (no Snowflake needed)
    def _add_ix(df):
        df["ix"] = (100 * df.mem_share / df.non_share).where(df.non_share > 0)
        return df

    res = dict(
        frm="2026-10-14", to="2026-10-21", generated="21 Oct 2026 10:00",
        kpis=dict(members=100, non_customers=900, mem_sales=1_500_000.0, non_sales=8_000_000.0,
                  mem_bills=300, non_bills=2500, abv_m=5000.0, abv_n=3200.0, ipb_m=2.4, ipb_n=2.0,
                  asp_m=2100.0, asp_n=1600.0, freq_m=3.0, freq_n=2.8, tot_sales=9_500_000.0,
                  mem_share_pct=15.8, mem_gmp=38.2, non_gmp=41.0, mem_cov=92.0, mem_gm=573000.0,
                  card_rev=37000.0, slab_disc=120000.0, n_emp_mem=20, emp_mem_sales=90000.0,
                  emp_card_paid=30000.0, emp_card_bills=25, n_empdisc_bills=40, empdisc_amt=60000.0,
                  n_dd_mob=5, n_win=10, n_onb=50, n_upg=2, n_trd=8, n_sup=20, unmapped_pct=4.0),
        insights=[("Basket", "Members spend ₹5,000 per bill vs ₹3,200 (+56%).", "Stretch baskets."),
                  ("Tiers", "Gold holds the largest share.", "Pitch upgrades.")],
        weekly=pd.DataFrame(dict(wk=pd.to_datetime(["2026-10-12", "2026-10-19"]),
                                 mem_v=[700000.0, 800000.0], tot_v=[4400000.0, 5100000.0])),
        enrol=pd.DataFrame(dict(wk=pd.to_datetime(["2026-10-12", "2026-10-19"]), new_members=[10, 12])),
        enrol_tier=pd.DataFrame(dict(wk=pd.to_datetime(["2026-10-12", "2026-10-19"]), rnk=[1, 2], n=[10, 12])),
        div=pd.DataFrame(dict(division=["MENS", "LADIES"], mem_v=[900000.0, 600000.0], mem_q=[400, 300],
                              mem_bills=[250, 180], non_v=[5000000.0, 3000000.0],
                              mem_share=[0.6, 0.4], non_share=[0.625, 0.375],
                              e_v=[100000.0, 100000.0], v_v=[300000.0, 200000.0],
                              p_v=[400000.0, 200000.0], sp_v=[100000.0, 100000.0])),
        subdiv=pd.DataFrame(dict(division=["MENS"], sub_division=["SHIRT"], mem_v=[500000.0], mem_q=[200],
                                 mem_bills=[150], non_v=[2000000.0], mem_share=[1.0], non_share=[1.0],
                                 e_v=[0.0], v_v=[100000.0], p_v=[300000.0], sp_v=[100000.0])),
        majcat=pd.DataFrame(dict(division=["MENS"], maj_cat=["FORMAL SHIRT"], mem_v=[500000.0], mem_q=[200],
                                 mem_bills=[150], non_v=[2000000.0], mem_share=[1.0], non_share=[1.0],
                                 e_v=[0.0], v_v=[100000.0], p_v=[300000.0], sp_v=[100000.0])),
        grid=pd.DataFrame(dict(grid=["E", "V", "P", "SP"], mem_v=[200000.0, 400000.0, 600000.0, 300000.0],
                               mem_q=[100, 200, 250, 100], mem_bills=[80, 150, 200, 90],
                               non_v=[2000000.0, 3000000.0, 2500000.0, 500000.0],
                               mem_share=[0.13, 0.27, 0.4, 0.2], non_share=[0.25, 0.375, 0.3125, 0.0625],
                               e_v=[200000.0, 0, 0, 0], v_v=[0, 400000.0, 0, 0],
                               p_v=[0, 0, 600000.0, 0], sp_v=[0, 0, 0, 300000.0])),
        size=pd.DataFrame(dict(division=["MENS", "MENS"], size=["M", "L"], mem_v=[400000.0, 500000.0], mem_q=[160, 200],
                               mem_bills=[120, 140], non_v=[1500000.0, 2000000.0],
                               mem_share=[0.44, 0.56], non_share=[0.43, 0.57],
                               e_v=[0.0, 0.0], v_v=[100000.0, 100000.0], p_v=[600000.0, 500000.0], sp_v=[200000.0, 200000.0])),
        colour=pd.DataFrame(dict(colour=["BLUE", "BLACK"], mem_v=[600000.0, 900000.0], mem_q=[250, 350],
                                 mem_bills=[180, 220], non_v=[2500000.0, 3500000.0],
                                 mem_share=[0.4, 0.6], non_share=[0.42, 0.58],
                                 e_v=[0.0, 0.0], v_v=[100000.0, 100000.0], p_v=[1000000.0, 900000.0], sp_v=[400000.0, 300000.0])),
        colour_div=pd.DataFrame(dict(division=["MENS"], colour=["BLUE"], mem_v=[600000.0], mem_q=[250],
                                     mem_bills=[180], non_v=[2500000.0], mem_share=[1.0], non_share=[1.0],
                                     e_v=[0.0], v_v=[100000.0], p_v=[1000000.0], sp_v=[400000.0])),
        grid_div=pd.DataFrame(dict(division=["MENS"], grid=["P"], mem_v=[600000.0], mem_q=[250],
                                   mem_bills=[200], non_v=[2500000.0], mem_share=[1.0], non_share=[1.0],
                                   e_v=[0.0], v_v=[0.0], p_v=[600000.0], sp_v=[0.0])),
        tier_tbl=pd.DataFrame(dict(members=[60, 30, 10], card_paid=[12000.0, 15000.0, 10000.0],
                                   mem_bills=[180, 90, 30], mem_sales=[900000.0, 450000.0, 150000.0],
                                   mem_abv=[5000.0, 5000.0, 5000.0], freq=[3.0, 3.0, 3.0],
                                   gmp=[36.0, 39.0, 42.0], slab_disc=[60000.0, 40000.0, 20000.0],
                                   card_rev=[12000.0, 15000.0, 10000.0], buyers=[60, 30, 10],
                                   e_pct=[0.2, 0.1, 0.05], v_pct=[0.3, 0.2, 0.15],
                                   p_pct=[0.35, 0.4, 0.4], sp_pct=[0.15, 0.3, 0.4]),
                             index=pd.Index(["Silver", "Gold", "Platinum"], name="tier")),
        gm_div=pd.DataFrame(dict(division=["MENS"], mem_v=[900000.0], mem_gm=[330000.0], mem_costed=[860000.0],
                                 non_v=[5000000.0], non_gm=[2050000.0], non_costed=[4800000.0])),
        card_rev=pd.DataFrame(dict(tier=["Silver", "Gold", "Platinum"], card_bills=[60, 30, 10],
                                   card_rev=[12000.0, 15000.0, 10000.0], buyers=[60, 30, 10])),
        members=pd.DataFrame(dict(mobile=["9800000001", "9800000003"], tier=["Silver", "Gold"],
                                  card_since=pd.to_datetime(["2026-10-01", "2026-10-02"]),
                                  card_paid=[200.0, 500.0], card_bills=[1, 1], mem_bills=[3, 2],
                                  mem_sales=[15000.0, 30000.0], mem_qty=[6, 10], mem_ev=[10000.0, 20000.0],
                                  last_dt=pd.to_datetime(["2026-10-05", "2026-10-15"]),
                                  store=["S001", "S002"])),
        emp_mem=pd.DataFrame(dict(mobile=["9800000002"], tier=["Gold"], card_since=pd.to_datetime(["2026-10-01"]),
                                  card_paid=[500.0], card_bills=[1], mem_bills=[2], mem_sales=[9000.0],
                                  mem_qty=[4], mem_ev=[5000.0], last_dt=pd.to_datetime(["2026-10-10"]),
                                  store=["S002"], ecode=["E001"], name=["Test Emp"], desig=["Exec"], loc=["HO"],
                                  status=["Active"])),
        emp_dd=pd.DataFrame(dict(d=[datetime.date(2026, 10, 10)], transactionid=["T1"], store=["S001"],
                                 mobile=["9800000002"], empid=["E001"], empdisc=[500.0])),
        campaigns=dict(winback=dict(file="a.csv", n=10, value=50000.0), onboard=dict(file="b.csv", n=50, value=200000.0),
                       upgrade=dict(file="c.csv", n=2, value=100000.0), tradeup=dict(file="d.csv", n=8, value=120000.0),
                       suppress=dict(file="e.csv", n=20, value=90000.0)),
        customers=pd.DataFrame(dict(
            mobile=["9800000001", "9800000003"], tier=["Silver", "Gold"],
            card_store=["S001", "S002"], card_since=pd.to_datetime(["2026-10-01", "2026-10-02"]),
            days_since=[15, 14], bef_visits_all=[0, 3], bef_last=[pd.NaT, pd.to_datetime("2026-01-15")],
            bef_visits_12m=[0, 1], bef_bills_12m=[0, 2], bef_val_12m=[0.0, 8000.0],
            bef_abv_12m=[pd.NA, 4000.0], cd_bills=[1, 1], cd_val=[3000.0, 5000.0],
            cd_abv=[3000.0, 5000.0], cd_disc=[60.0, 250.0], aft_visits=[2, 3], aft_bills=[3, 4],
            aft_val=[9000.0, 12000.0], aft_abv=[3000.0, 3000.0], aft_disc=[180.0, 600.0],
            ttl_disc=[240.0, 850.0], card_fee=[200, 500], disc_pct_fee=[1.2, 1.7],
            fee_recovered=["<50%", "FULLY"], is_emp=["N", "Y"], status=["", "Active"],
            high_disc=["N", "Y"], rare_shopper=["N", "Y"], new_cust=["Y", "N"],
            flag=["NEW CUST", "HIGH DISC, RARE SHOPPER"])))
    p = export({**res,
                "div": _add_ix(res["div"]), "subdiv": _add_ix(res["subdiv"]),
                "majcat": _add_ix(res["majcat"]), "grid": _add_ix(res["grid"]),
                "size": _add_ix(res["size"]), "colour": _add_ix(res["colour"]),
                "colour_div": _add_ix(res["colour_div"]), "grid_div": _add_ix(res["grid_div"])},
               os.path.join(OUT, "_selftest.xlsx"))
    print("self-test OK ->", p)
